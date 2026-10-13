"""Exact matching, renderer parity, bounded ambiguity and real PDF extraction."""

import json
import shutil
from copy import deepcopy

import pytest
from pptx import Presentation
from pptx.util import Inches

from src.agents.renderer import Renderer
from src.agents.render_fit_audit import render_fit_audit_node
from src.layout.render_fit import audit_rendered_text, evaluate_rendered_boxes
from utils.text_formatting import editable_text_parts, tokenize_formatting, tokenize_styling


def _element(text, sid="s", x=1, y=1, kind="section_title", width=4, height=2):
    return {"id": sid + ("_title" if kind == "section_title" else "_text"),
            "section_id": sid, "zone_id": "main", "type": kind,
            "section_title" if kind == "section_title" else "content": text,
            "x": x, "y": y, "width": width, "height": height}


def _pdf(groups):
    """Each group is a merged PDF block; tuples are independent text lines."""
    raw, blocks = {"blocks": []}, []
    for index, group in enumerate(groups):
        lines = []
        for text, x, y in group:
            chars = [{"c": char, "bbox": [x + offset * 5, y, x + (offset + 1) * 5, y + 12]}
                     for offset, char in enumerate(text)]
            lines.append({"spans": [{"chars": chars}]})
        boxes = [char["bbox"] for line in lines for span in line["spans"] for char in span["chars"]]
        box = [min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes)]
        raw["blocks"].append({"type": 0, "number": index, "lines": lines})
        blocks.append((*box, "\n".join(line[0] for line in group), index, 0))
    return blocks, raw


def _audit(layout, groups, **kwargs):
    blocks, raw = _pdf(groups)
    return evaluate_rendered_boxes(layout, blocks, raw_pdf=raw, poster_width=18, poster_height=12, **kwargs)


@pytest.mark.parametrize("text,visible", [
    ("**Bold** and *italic*", "Bold and italic"),
    ("<color:#137a7f>HAGS</color> results", "HAGS results"),
    ("<color:#abcdef>**literal**</color>", "**literal**"),
    (r"\*literal\* \\alpha a*b", r"*literal* \alpha a*b"),
    ("<color:#123456>unclosed", "<color:#123456>unclosed"),
    ("<color:red>literal</color>", "<color:red>literal</color>"),
    # Legacy grammar interprets adjacent ** as empty italic when bold fails.
    ("**unclosed", "unclosed"),
    ("**two\nlines**", "two\nlines"),
    ("<color:#abcdef> </color>remaining", "remaining"),
    (r"H_2O and x^{2}", "H2O and x2"),
])
def test_shared_parser_preserves_legacy_visible_runs(text, visible):
    renderer = Renderer()
    assert renderer._tokenize_styling(text) == tokenize_styling(text)
    assert renderer._tokenize_formatting(text) == tokenize_formatting(text)
    assert "".join(run["text"] for run in tokenize_formatting(text)) == visible
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(0, 0, Inches(10), Inches(2)).text_frame.paragraphs[0]
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    renderer._parse_and_add_runs(paragraph, text, "Arial", Pt(40), RGBColor(0, 0, 0))
    assert paragraph.text == visible


def test_style_attributes_are_unchanged_and_color_content_is_not_recursive():
    assert tokenize_styling("**b** *i* <color:#123456>**c**</color>") == [
        {"text": "b", "bold": True, "italic": False, "color": None},
        {"text": " ", "bold": False, "italic": False, "color": None},
        {"text": "i", "bold": False, "italic": True, "color": None},
        {"text": " ", "bold": False, "italic": False, "color": None},
        {"text": "**c**", "bold": True, "italic": False, "color": "#123456"},
    ]


def test_styled_body_matches_visible_text_not_color_tag_letters():
    element = _element("<color:#137a7f>HAGS</color> has grounded results.", kind="text")
    report = _audit([element], [[("HAGS has grounded results.", 79, 78)]])
    assert report["status"] == "passed"
    assert report["matched_text_boxes"][0]["match_method"] == "exact_block"
    assert report["matching_diagnostics"]["fallback_elements"] == 0
    assert "grounded" not in json.dumps(report)


@pytest.mark.parametrize("columns", [2, 3])
@pytest.mark.parametrize("wrapped", [False, True])
def test_merged_parallel_titles_have_disjoint_complete_character_evidence(columns, wrapped):
    layout, lines = [], []
    for index in range(columns):
        x = 1 + 5 * index
        layout.append(_element(f"Scientific results {index}", sid=str(index), x=x))
        lines.append(("Scientific" if wrapped else f"Scientific results {index}", x * 72 + 7, 78))
        if wrapped:
            lines.append((f"results {index}", x * 72 + 7, 96))
    report = _audit(layout, [lines])
    assert report["status"] == "passed", report
    assert report["matched_text_count"] == columns
    assert all(item["match_method"] == "exact_fragments" for item in report["matched_text_boxes"])
    assert all(item["fragment_count"] == (2 if wrapped else 1) for item in report["matched_text_boxes"])


def test_body_can_span_independent_paragraph_blocks_and_style_spans():
    element = _element("**First** grounded statement.\nSecond verified statement.", kind="text")
    report = _audit([element], [[("First grounded statement.", 79, 78)], [("Second verified statement.", 79, 98)]])
    assert report["status"] == "passed"
    assert report["matched_text_boxes"][0]["fragment_count"] == 2


def test_character_subrange_handles_two_independent_titles_in_one_pdf_line():
    first = _element("Alpha results", width=2)
    second = _element("Beta results", sid="b", x=2.1, width=2)
    # PDF happens to emit a single line. Disjoint subranges, not shared block.
    report = _audit([first, second], [[("Alpha results   Beta results", 79, 78)]])
    assert report["matched_text_count"] == 2


@pytest.mark.parametrize("case", ["missing", "reordered", "prefix_only", "unrelated_suffix"])
def test_incomplete_or_reordered_body_is_not_full_coverage(case):
    element = _element("First grounded statement. Second verified statement.", kind="text")
    groups = {
        "missing": [[("First grounded statement.", 79, 78)]],
        "reordered": [[("Second verified statement.", 79, 78), ("First grounded statement.", 79, 98)]],
        "prefix_only": [[("First grounded statement. Second verified", 79, 78)]],
        "unrelated_suffix": [[("First grounded statement. Second invented statement.", 79, 78)]],
    }
    report = _audit([element], groups[case])
    assert report["status"] == "partial"
    assert report["matched_text_count"] == 0
    assert report["passed"] and not report["conclusive"]


def test_ambiguous_fragments_are_not_resolved_by_nearest_origin():
    element = _element("Repeated results")
    report = _audit([element], [[("Repeated results", 79, 78), ("Repeated results", 79, 98)]])
    assert report["status"] == "partial"
    assert report["unmatched_text_reasons"][0]["reason"] == "ambiguous_fragments"


def test_same_pdf_characters_cannot_satisfy_two_elements():
    first = _element("Shared results")
    second = {**first, "id": "other_title"}
    report = _audit([first, second], [[("Shared results", 79, 78)]])
    assert report["matched_text_count"] == 0
    assert {item["reason"] for item in report["unmatched_text_reasons"]} == {"shared_pdf_evidence"}


def test_matched_union_keeps_overflowing_last_line_and_collision():
    title = _element("First results second line", height=0.3)
    other = _element("Other module", sid="b", y=1.4)
    report = _audit([title, other], [[("First results", 79, 78), ("second line", 79, 105), ("Other module", 79, 140)]])
    assert report["status"] == "failed"
    assert report["matched_text_boxes"][0]["rendered_box_inches"][3] == 1.625
    assert {item["code"] for item in report["issues"]} >= {"RENDER_TEXT_BOX_OVERFLOW", "RENDER_CROSS_SECTION_COLLISION"}


def test_known_failure_survives_another_unmatched_element():
    first = _element("Confirmed overflow", height=0.1)
    other = _element("Missing text", sid="b", x=10)
    report = _audit([first, other], [[("Confirmed overflow", 79, 78)]])
    assert report["status"] == "failed" and not report["passed"]
    assert not report["conclusive"]


@pytest.mark.parametrize("limit", ["MAX_PDF_CHARACTERS", "MAX_PDF_LINES", "MAX_SEARCH_STEPS", "MAX_PATH_LINES", "MAX_ACTIVE_PATHS", "MAX_MATCH_CHARACTERS"])
def test_fragment_limits_fail_inconclusively(monkeypatch, limit):
    monkeypatch.setattr(f"src.layout.render_text_matching.{limit}", 0)
    report = _audit([_element("First results second line")], [[("First results", 79, 78), ("second line", 79, 96), ("Independent title", 600, 78)]])
    assert report["status"] == "partial" and report["passed"]
    assert "budget_exhausted" in report["unmatched_text_reasons"][0]["reason"]


def test_raw_pdf_is_lazy_and_called_once():
    layout = [_element("Exact results")]
    blocks, raw = _pdf([[("Exact results", 79, 78)]])
    calls = []
    def extract():
        calls.append(True)
        return raw
    report = evaluate_rendered_boxes(layout, blocks, raw_pdf=extract, poster_width=18, poster_height=12)
    assert report["status"] == "passed" and not calls
    layout += [_element("Missing results", sid="b", x=10), _element("Another missing title", sid="c", x=15)]
    report = evaluate_rendered_boxes(layout, blocks, raw_pdf=extract, poster_width=18, poster_height=12)
    assert report["status"] == "partial" and len(calls) == 1


def test_fine_extraction_failure_keeps_confirmed_block_failure():
    blocks, _ = _pdf([[("Exact results", 79, 78)]])
    def unavailable():
        raise RuntimeError("PDF extraction unavailable")
    report = evaluate_rendered_boxes([_element("Exact results", height=0.01), _element("Missing results", sid="b")],
                                    blocks, raw_pdf=unavailable, poster_width=18, poster_height=12)
    assert report["status"] == "failed"
    assert report["unmatched_text_reasons"][0]["reason"] == "pdf_fragment_extraction_unavailable"


def test_display_formula_images_are_explicitly_outside_editable_audit_scope():
    text = "Grounded editable explanation.\nε(α) = α/2"
    parts, count = editable_text_parts(text)
    assert parts == ["Grounded editable explanation."] and count == 1
    report = _audit([_element(text, kind="text")], [[("Grounded editable explanation.", 79, 78)]])
    assert report["matched_text_count"] == 1
    assert report["excluded_noneditable_content"] == [{"element_id": "s_text", "kind": "display_formula_image", "count": 1, "text_verified": False}]


def test_title_markup_is_literal_unlike_body_markup():
    report = _audit([_element("<color:#123456>Results</color>")], [[("Results", 79, 78)]])
    assert report["status"] == "partial"


def test_formula_prefix_is_still_editable_and_requires_matching():
    parts, count = editable_text_parts("The grounded formula is ε(α) = α/2")
    assert parts == ["The grounded formula is"] and count == 1
    report = _audit([_element("The grounded formula is ε(α) = α/2", kind="text")], [[("Unrelated text", 79, 78)]])
    assert report["status"] == "partial"


@pytest.mark.parametrize("suffix", [" extra invented words", "invented"])
def test_fragment_match_does_not_hide_extra_words_or_partial_word(suffix):
    report = _audit([_element("Scientific results")], [[("Scientific results" + suffix, 79, 78)]])
    assert report["status"] == "partial" and report["matched_text_count"] == 0


def test_partial_terminal_summary_and_reason_do_not_leak_scientific_text(monkeypatch, tmp_path):
    report = _audit([_element("Missing scientific title")], [[("Unrelated text", 79, 78)]])
    monkeypatch.setattr("src.agents.render_fit_audit.audit_rendered_text", lambda _: report)
    logs = []
    monkeypatch.setattr("src.agents.render_fit_audit.log_agent_warning", lambda _, message: logs.append(message))
    render_fit_audit_node({"generation_route": "optimization_family", "output_dir": str(tmp_path)})
    assert any("RenderFit partial | matched 0/1 | titles 0/1 | body 0/0" in line for line in logs)
    assert any("s_title | exact_text_not_found" in line for line in logs)
    assert not any("scientific" in line for line in logs)


def test_unavailable_reports_explicit_inconclusive_coverage(tmp_path):
    report = audit_rendered_text({"styled_layout": [_element("Missing results")], "current_candidate": {"pptx_path": str(tmp_path / "missing.pptx")}})
    assert report["status"] == "unavailable" and not report["conclusive"]
    assert report["expected_text_count"] == 1
    assert report["matching_policy_version"] == "editable_exact_fragments_v1"


def test_real_libreoffice_parallel_wrapped_titles_and_styled_body(tmp_path, monkeypatch):
    if not shutil.which("libreoffice"):
        pytest.skip("LibreOffice not installed")
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(18), Inches(12)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    renderer = Renderer()
    renderer.styling_interfaces = {"font_sizes": {"section_title": 64, "body_text": 40}, "line_spacing": 1.0}
    layout = []
    for index in range(3):
        title = _element(f"Scientific results and evaluation {index}", sid=str(index), x=1 + 5.5 * index, height=3)
        title.update(font_size=64, font_family="Arial", font_weight="bold")
        renderer._render_section_title(slide, title, {"generation_route": "optimization_family"})
        layout.append(title)
        body = _element("<color:#137a7f>HAGS</color> yields **grounded** results.\nSecond verified paragraph.", sid=str(index), x=title["x"], y=5, kind="text", height=4)
        renderer._add_enhanced_text(slide, body["content"], Inches(body["x"]), Inches(body["y"]), Inches(4), Inches(4), body)
        layout.append(body)
    pptx = tmp_path / "parallel.pptx"
    presentation.save(pptx)
    import hashlib
    original = hashlib.sha256(pptx.read_bytes()).hexdigest()
    state = {"styled_layout": layout, "poster_width": 18, "poster_height": 12, "current_candidate": {"pptx_path": str(pptx)}}
    before = deepcopy(state)
    report = audit_rendered_text(state)
    assert report["matched_text_count"] == 6, report
    assert report["text_coverage"]["fraction"] == 1
    assert report["title_gap_coverage"] == {"expected": 3, "measured": 3}
    assert state == before
    assert original == hashlib.sha256(pptx.read_bytes()).hexdigest()
    assert not list(tmp_path.glob("*.pdf"))
    def missing_composition(*args, **kwargs):
        raise RuntimeError("Optional composition measurement unavailable")
    monkeypatch.setattr("src.layout.whitespace.measure_composition_whitespace", missing_composition)
    state["styled_layout"][0]["height"] = 0.01
    failed = audit_rendered_text(state)
    assert failed["status"] == "failed" and not failed["passed"]
    assert failed["composition_measurement_status"] == "unavailable"
    assert any(issue["code"] == "RENDER_TEXT_BOX_OVERFLOW" for issue in failed["issues"])

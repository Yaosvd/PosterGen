"""No-model coverage for accepted bundle and visual skeleton contracts."""

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from src.agents.layout_agent import LayoutAgent
from src.layout.families import LayoutArchitect, LayoutFamilyRegistry
from src.templates.library import LoopStorage, TemplateDistiller, TemplateLibrary
from src.templates.preview import generate_family_previews
from src.templates.transfer import assess_template_transfer
from src.workflow.campaign_runner import (
    _reconcile_accepted_bundles,
    _discard_uncommitted_bundles,
)


def _candidate(tmp_path):
    source = tmp_path / "candidate"
    source.mkdir()
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(source / "poster.pptx")
    Image.new("RGB", (540, 360), "white").save(source / "poster.png")
    candidate_id = "campaign_i000001_example"
    state = {
        "poster_name": "paper",
        "poster_width": 54,
        "poster_height": 36,
        "current_candidate": {
            "candidate_id": candidate_id,
            "pptx_path": str(source / "poster.pptx"),
            "png_path": str(source / "poster.png"),
        },
        "styled_layout": [
            {"type": "section_container", "x": 2, "y": 9, "width": 24,
             "height": 10, "zone_id": "left", "text": "SECRET SCIENCE"},
            {"type": "visual", "x": 3, "y": 12, "width": 10,
             "height": 5, "zone_id": "left"},
        ],
        "layout_spec": {
            "family_id": "asymmetric_two_column",
            "primary_anchor_zone": "evidence_field",
            "zones": [
                {"zone_id": "narrative_rail", "x": 0, "y": 0, "width": .31,
                 "height": 1, "role": "support", "capacity": 4, "section_ids": ["private"]},
                {"zone_id": "evidence_field", "x": .34, "y": 0, "width": .66,
                 "height": 1, "role": "hero", "capacity": 5, "section_ids": []},
            ],
        },
        "iteration_context": {
            "campaign_id": "campaign-source",
            "batch_id": "batch-3",
            "iteration": 7,
            "lineage_id": "lineage-source",
            "portfolio_mode": "improve",
            "parent_candidate_id": "parent-candidate",
        },
        "loop_state": {"paper_id": "paper-hash"},
    }
    return state


def test_accepted_bundle_has_one_poster_and_editable_paper_free_template(tmp_path):
    state = _candidate(tmp_path)
    report = {"archive_eligible": True, "overall_score": 89,
              "dimension_scores": {}, "candidate_id": state["current_candidate"]["candidate_id"]}
    template = TemplateDistiller().distill(
        state, report, {"profile_key": "balanced:medium", "visual_count": 2,
                        "section_count": 1, "content_density": "medium"}
    )
    storage = LoopStorage(root=tmp_path / "store")
    archive = storage.prepare_accepted_bundle("session", state, report, template)
    assert not Path(archive["archive_dir"]).exists()
    storage.publish_accepted_bundle(archive)
    library = TemplateLibrary(storage.root)
    assert library.register(template)["registered"]
    assert not library.register(template)["registered"]
    files = {path.name for path in Path(archive["archive_dir"]).iterdir()}
    assert {"poster.pptx", "poster.png", "template.json", "template.pptx",
            "template.png", "critic_report.json"} <= files
    assert json.loads(Path(archive["template_json_path"]).read_text())["source_candidate_id"] == archive["candidate_id"]
    saved_template = json.loads(Path(archive["template_json_path"]).read_text())
    assert saved_template["template_evaluation"]["rubric_version"] == "structural_template_v1"
    assert saved_template["template_design_score"] == saved_template["template_evaluation"]["score"]
    assert saved_template["source_campaign_id"] == "campaign-source"
    assert saved_template["source_batch_id"] == "batch-3"
    assert saved_template["source_iteration"] == 7
    assert saved_template["source_lineage_id"] == "lineage-source"
    assert saved_template["source_mode"] == "improve"
    assert saved_template["source_parent_candidate_id"] == "parent-candidate"
    assert saved_template["lineage_provenance_status"] == "complete"
    provenance = json.loads(
        (Path(archive["archive_dir"]) / "provenance.json").read_text()
    )
    assert provenance["schema_version"] == 3
    preview = Presentation(archive["template_pptx_path"])
    assert preview.slide_width / preview.slide_height == pytest.approx(1.5, rel=.001)
    assert len(preview.slides[0].shapes) >= 4
    assert "SECRET SCIENCE" not in " ".join(shape.text for shape in preview.slides[0].shapes)
    with Image.open(archive["template_png_path"]) as image:
        assert image.width / image.height == pytest.approx(1.5, rel=.002)


def test_missing_poster_artifact_never_publishes_bundle(tmp_path):
    state = _candidate(tmp_path)
    Path(state["current_candidate"]["png_path"]).unlink()
    report = {"archive_eligible": True, "overall_score": 89}
    template = {"source_candidate_id": state["current_candidate"]["candidate_id"]}
    storage = LoopStorage(root=tmp_path / "store")
    with pytest.raises(FileNotFoundError):
        storage.prepare_accepted_bundle("session", state, report, template)
    assert not list((storage.root / "accepted_posters" / ".pending").glob("*/template.json"))


def test_all_registered_families_have_editable_png_pptx_skeletons(tmp_path):
    paths = generate_family_previews(tmp_path)
    assert len(paths) == 24
    for family_id in LayoutFamilyRegistry().ids():
        assert (tmp_path / f"{family_id}.png").is_file()
        deck = Presentation(tmp_path / f"{family_id}.pptx")
        assert len(deck.slides[0].shapes) >= 2
        assert all(shape.has_text_frame for shape in deck.slides[0].shapes)


def test_strict_validation_preserves_topology_and_flags_mutation(tmp_path):
    state = _candidate(tmp_path)
    topology = TemplateDistiller().distill(
        state, {"overall_score": 89}, {"profile_key": "balanced:medium"}
    )["zone_topology"]
    state["template_validation_topology"] = topology
    state["validation_template_id"] = "T1"
    state["story_board"] = {"spatial_content_plan": {"sections": [
        {"section_id": "s1", "claim_ids": ["c1"], "visual_assets": []}
    ]}}
    spec = LayoutArchitect().build(state, "asymmetric_two_column")
    state["layout_spec"] = spec
    assert assess_template_transfer(state)["passed"] is True
    spec["zones"][0]["width"] -= .05
    assert assess_template_transfer(state)["passed"] is False


def test_replay_publishes_and_registers_only_committed_bundle(tmp_path):
    state = _candidate(tmp_path)
    report = {"archive_eligible": True, "overall_score": 89,
              "candidate_id": state["current_candidate"]["candidate_id"]}
    template = TemplateDistiller().distill(
        state, report, {"profile_key": "balanced:medium"}
    )
    store = LoopStorage(root=tmp_path / "store")
    library = TemplateLibrary(store.root)
    owner = SimpleNamespace(storage=store, library=library)
    handle = SimpleNamespace(path=tmp_path / "campaign")
    attempt = handle.path / "iterations" / "000001" / "attempts" / "attempt_001"
    attempt.mkdir(parents=True)
    archive = store.prepare_accepted_bundle("s", state, report, template)
    (attempt / "outcome.json").write_text(json.dumps({
        "accepted": True, "archive": archive,
        "template_registration": {"template": template},
    }))
    marker = attempt.parents[1] / "COMPLETED.json"
    marker.write_text(json.dumps({"attempt_id": "attempt_001"}))
    _reconcile_accepted_bundles(owner, handle)
    library.validation_queue_path.unlink()
    _reconcile_accepted_bundles(owner, handle)
    assert library.validation_queue_path.is_file()
    assert Path(archive["template_pptx_path"]).is_file()
    assert len(library._load()["templates"]) == 1


def test_uncommitted_bundle_is_discarded_on_recovery(tmp_path):
    state = _candidate(tmp_path)
    report = {"archive_eligible": True, "overall_score": 89}
    template = TemplateDistiller().distill(
        state, report, {"profile_key": "balanced:medium"}
    )
    store = LoopStorage(root=tmp_path / "store")
    owner = SimpleNamespace(storage=store)
    handle = SimpleNamespace(path=tmp_path / "campaign")
    attempt = handle.path / "iterations" / "000001" / "attempts" / "attempt_001"
    attempt.mkdir(parents=True)
    (attempt / "status.json").write_text(json.dumps({
        "candidate_id": state["current_candidate"]["candidate_id"]
    }))
    archive = store.prepare_accepted_bundle("s", state, report, template)
    _discard_uncommitted_bundles(owner, handle)
    assert not Path(archive["archive_dir"]).exists()
    assert not (store.root / "accepted_posters" / ".pending" / archive["candidate_id"]).exists()


def test_validation_topology_does_not_leak_into_next_exploration(tmp_path):
    state = _candidate(tmp_path)
    state["story_board"] = {"spatial_content_plan": {"sections": [
        {"section_id": "s1", "claim_ids": ["c1"], "visual_assets": []}
    ]}}
    state["template_validation_topology"] = {
        "family_id": "asymmetric_two_column",
        "primary_anchor_zone": "evidence_field",
        "zones": [{key: zone[key] for key in
                   ("zone_id", "x", "y", "width", "height", "role", "capacity")}
                  for zone in state["layout_spec"]["zones"]],
    }
    state["validation_template_id"] = "T1"
    state["iteration_context"] = {"layout_family": "asymmetric_two_column", "random_seed": 1}
    locked = LayoutAgent._ensure_layout_spec(state, state["story_board"])
    assert locked["validation_template_id"] == "T1"
    state["template_validation_topology"] = {}
    state["validation_template_id"] = None
    state["iteration_context"] = {"layout_family": "asymmetric_two_column", "random_seed": 2}
    fresh = LayoutAgent._ensure_layout_spec(state, state["story_board"])
    assert fresh is not locked
    assert "validation_template_id" not in fresh
    assert fresh["iteration_random_seed"] == 2


def test_validation_evidence_saves_poster_without_new_template(tmp_path):
    state = _candidate(tmp_path)
    state["iteration_context"] = {
        "portfolio_mode": "validate",
        "validation_target": "cross_paper",
    }
    report = {
        "archive_eligible": True,
        "overall_score": 94,
        "candidate_id": state["current_candidate"]["candidate_id"],
        "task_evaluation": {"passed": True, "reasons": []},
    }
    storage = LoopStorage(root=tmp_path / "store")
    evidence = storage.prepare_validation_evidence(
        "session",
        state,
        report,
        template_id="template",
        template_fingerprint="fingerprint",
        paper_id="paper-hash",
    )
    assert evidence["saved"] is True
    assert not Path(evidence["archive_dir"]).exists()
    storage.publish_validation_evidence(evidence)
    saved = Path(evidence["archive_dir"])
    files = {path.name for path in saved.iterdir()}
    assert {
        "poster.pptx",
        "poster.png",
        "critic_report.json",
        "task_evaluation.json",
        "story_board.json",
        "styled_layout.json",
        "provenance.json",
    } <= files
    assert not any(name.startswith("template.") for name in files)
    assert not (storage.root / "accepted_posters").exists()
    provenance = json.loads((saved / "provenance.json").read_text())
    assert provenance["validated_template_id"] == "template"
    assert provenance["artifact_type"] == "strict_template_validation_evidence"


def test_uncommitted_validation_evidence_is_discarded_on_recovery(tmp_path):
    state = _candidate(tmp_path)
    report = {
        "archive_eligible": True,
        "overall_score": 94,
        "task_evaluation": {"passed": True},
    }
    storage = LoopStorage(root=tmp_path / "store")
    owner = SimpleNamespace(storage=storage)
    handle = SimpleNamespace(path=tmp_path / "campaign")
    attempt = handle.path / "iterations" / "000001" / "attempts" / "attempt_001"
    attempt.mkdir(parents=True)
    (attempt / "status.json").write_text(json.dumps({
        "candidate_id": state["current_candidate"]["candidate_id"]
    }))
    evidence = storage.prepare_validation_evidence(
        "session",
        state,
        report,
        template_id="template",
        template_fingerprint="fingerprint",
        paper_id="paper-hash",
    )
    _discard_uncommitted_bundles(owner, handle)
    assert not Path(evidence["archive_dir"]).exists()
    assert not (
        storage.root
        / "validation_evidence"
        / ".pending"
        / evidence["candidate_id"]
    ).exists()

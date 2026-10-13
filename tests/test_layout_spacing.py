"""Title/render consistency, directional whitespace and bounded refinement."""

from copy import deepcopy
import shutil
import subprocess

import pytest
from pptx import Presentation
from pptx.util import Inches

from src.agents.layout_agent import LayoutAgent
from src.agents.renderer import Renderer
from src.agents.render_fit_audit import _correct_title_fit
from src.config.poster_config import load_config
from src.config.critic_config import load_critic_config
from src.critic.base import CriticContext
from src.critic.section_fit import SectionFitEvaluator
from src.layout.capacity_repair import repair_improves_fit
from src.layout.composition_refinement import boundary_options, refine_composition
from src.layout.feasibility import validate_intra_section_layout
from src.layout.measurement_context import measurement_scope
from src.layout.render_fit import evaluate_rendered_boxes
from src.layout.title_metrics import measure_loop_title, title_policy
from src.layout.whitespace import measure_composition_whitespace, module_adjacencies, no_displaced_whitespace


@pytest.mark.parametrize("title,width", [
    ("Results", 13.851111),
    ("Ablation Studies and Feature Utility", 13.851111),
    ("Multi-Modal Data Integration and Experimental", 13.851111),
    ("Performance Evaluation: HAGS vs. Baselines", 13.851111),
    ("A long title with explicit\nline breaks", 8.0),
])
def test_title_cache_matches_disabled_and_uncached(title, width):
    config = load_config()
    baseline = measure_loop_title(title, width, "Helvetica Neue", 64, config)
    with measurement_scope() as context:
        assert measure_loop_title(title, width, "Helvetica Neue", 64, config) == baseline
        assert measure_loop_title(title, width, "Helvetica Neue", 64, config) == baseline
        assert context.stats["line"]["hits"] > 0
    assert not context.cache
    with measurement_scope({"enabled": False}):
        assert measure_loop_title(title, width, "Helvetica Neue", 64, config) == baseline


def test_title_policy_change_cannot_reuse_old_height():
    config = load_config()
    with measurement_scope():
        old = measure_loop_title("Methods and Results", 6, "Arial", 64, config)
        changed = deepcopy(config)
        changed["text_measurement"]["loop_section_title"]["line_advance"] = 1.5
        new = measure_loop_title("Methods and Results", 6, "Arial", 64, changed)
        assert new["height"] > old["height"]
        changed["text_measurement"]["loop_section_title"]["height_reserve_inches"] = 0.4
        assert measure_loop_title("Methods and Results", 6, "Arial", 64, changed)["height"] > new["height"]


@pytest.mark.parametrize("value", [-1, float("nan"), 10])
def test_invalid_title_policy_fails(value):
    config = load_config()
    config["text_measurement"]["loop_section_title"]["line_advance"] = value
    with pytest.raises(ValueError):
        title_policy(config)


def test_deterministic_final_font_preset_is_measured_before_compile():
    agent = LayoutAgent()
    state = {"generation_route": "optimization_family", "style_overrides": {"font_family": "normalized"}}
    metrics = agent._measure_section_title("Methods and Results", 10, state)
    assert metrics["font_name"] == "Arial"


@pytest.mark.parametrize("title", ["Results", "Ablation Studies and Feature Utility",
                                  "Multi-Modal Data Integration and Experimental",
                                  "Performance Evaluation: HAGS vs. Baselines"])
def test_real_libreoffice_title_matches_height_and_spacing(tmp_path, title):
    if not shutil.which("libreoffice"):
        pytest.skip("LibreOffice not installed")
    import pymupdf
    config = load_config()
    metrics = measure_loop_title(title, 13.851111, "Helvetica Neue", 64, config)
    element = {"type": "section_title", "section_id": "s", "id": "s_title", "zone_id": "main",
               "section_title": title, "x": 1, "y": 1, "width": 13.851111,
               "font_size": 64, "font_family": "Helvetica Neue", "font_weight": "bold", **metrics}
    visual = {"type": "visual", "section_id": "s", "x": 1, "y": 1 + metrics["height"] + 0.4,
              "width": 12, "height": 2}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(18), Inches(12)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    renderer = Renderer()
    renderer.styling_interfaces = {"font_sizes": {"section_title": 64}}
    renderer._render_section_title(slide, element, {"generation_route": "optimization_family"})
    paragraph = slide.shapes[0].text_frame.paragraphs[0]
    assert paragraph.line_spacing.pt == pytest.approx(76.8)
    assert paragraph.space_before.pt == paragraph.space_after.pt == 0
    pptx = tmp_path / "titles.pptx"
    prs.save(pptx)
    subprocess.run([shutil.which("libreoffice"), f"-env:UserInstallation={(tmp_path / 'profile').as_uri()}",
                    "--headless", "--convert-to", "pdf:impress_pdf_Export", "--outdir", str(tmp_path), str(pptx)],
                   check=True, timeout=120, capture_output=True)
    with pymupdf.open(tmp_path / "titles.pdf") as document:
        page = document[0]
        report = evaluate_rendered_boxes([element, visual], page.get_text("blocks"), poster_width=18, poster_height=12)
        page.get_pixmap(matrix=pymupdf.Matrix(0.5, 0.5)).save(tmp_path / "titles.png")
    assert report["matched_text_count"] == 1
    assert report["status"] == "passed", report
    assert report["internal_title_gaps"][0]["gap_inches"] >= 0.2


def test_classic_title_keeps_implicit_spacing():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    renderer = Renderer()
    renderer.styling_interfaces = {"font_sizes": {"section_title": 64}}
    renderer._render_section_title(slide, {"x": 1, "y": 1, "width": 10, "height": 3,
                                          "section_title": "Results"}, {"generation_route": "standalone_classic"})
    assert slide.shapes[0].text_frame.paragraphs[0].line_spacing is None


def _title_layout():
    return [{"type": "section_title", "id": "s_title", "section_id": "s", "zone_id": "top",
             "section_title": "Results", "x": 1, "y": 1, "width": 4, "height": 1},
            {"type": "visual", "section_id": "s", "x": 1, "y": 2.4, "width": 4, "height": 2}]


def test_actual_title_gap_catches_crowding_without_large_overlap():
    report = evaluate_rendered_boxes(_title_layout(), [(79, 85, 290, 172, "Results", 0, 0)], poster_width=8, poster_height=8)
    assert any(item["code"] == "RENDER_TITLE_CONTENT_GAP" for item in report["issues"])


def test_short_repeated_titles_match_by_position_and_ambiguous_stays_partial():
    first = _title_layout()[0]
    second = {**first, "id": "t_title", "section_id": "t", "y": 5}
    blocks = [(79, 85, 290, 130, "Results", 0, 0), (79, 373, 290, 420, "Results", 1, 0)]
    report = evaluate_rendered_boxes([first, second], blocks, poster_width=8, poster_height=8)
    assert report["matched_text_count"] == 2
    report = evaluate_rendered_boxes([first], blocks[:1] * 2, poster_width=8, poster_height=8)
    assert report["status"] == "partial" and not report["conclusive"]


@pytest.mark.parametrize("mode", ["validate", "explore", "improve"])
def test_correction_does_not_touch_locked_or_exact_fallback(mode):
    state = {"iteration_context": {"portfolio_mode": mode}, "improve_safe_fallback": True}
    original = deepcopy(state)
    assert _correct_title_fit(state, {"issues": []}) is None
    assert state == original


def _container(sid, zone, x, y, width, height):
    return {"type": "section_container", "section_id": sid, "zone_id": zone,
            "x": x, "y": y, "width": width, "height": height, "planned_zone_height": height + 0.1}


def _whitespace(layout):
    return measure_composition_whitespace({"zones": []}, layout, poster_width=12, poster_height=12,
                                         poster_margin=1, title_height_fraction=0.1)


def test_close_horizontal_neighbor_does_not_mask_large_vertical_gap():
    layout = [_container("a", "top", 1, 2, 4, 2), _container("b", "lower", 1, 7, 4, 3),
              _container("c", "side", 5.2, 2, 4, 8)]
    result = _whitespace(layout)
    assert result["largest_vertical_gap_inches"] == 3
    assert result["largest_horizontal_gap_inches"] == 0.2
    assert result["largest_cross_zone_gap_inches"] == 3


def test_occluded_non_neighbor_is_not_counted():
    layout = [_container("a", "a", 1, 2, 4, 1), _container("b", "b", 1, 4, 4, 1),
              _container("c", "c", 1, 6, 4, 1)]
    pairs = {(item["after_section_id"], item["before_section_id"]) for item in module_adjacencies(layout)}
    assert ("a", "c") not in pairs
    assert ("a", "b") in pairs and ("b", "c") in pairs


def test_bottom_gap_is_extra_inside_body_not_outer_margin():
    result = _whitespace([_container("a", "a", 1, 2, 10, 7)])
    assert result["mean_terminal_bottom_gap_inches"] == 2
    assert result["full_width_bottom_gap_inches"] == 2
    assert result["bottom_edge_coverage_fraction"] == 1


def test_moving_gap_to_bottom_is_not_an_improvement():
    before = {"schema_version": 2, "largest_vertical_gap_inches": 2,
              "mean_terminal_bottom_gap_inches": 1}
    after = {"schema_version": 2, "largest_vertical_gap_inches": 1,
             "mean_terminal_bottom_gap_inches": 2}
    assert not no_displaced_whitespace(before, after)
    assert no_displaced_whitespace({}, before)  # historical unknowns are not zero measurements


def test_local_repair_gain_allowed_when_worst_region_is_unchanged():
    comparison = {"zone_utilization_before": {"a": 0.85, "b": 0.90},
                  "zone_utilization_after": {"a": 0.85, "b": 0.94},
                  "worst_utilization_before": 0.85, "worst_utilization_after": 0.85,
                  "composition_before": {}, "composition_after": {},
                  "bottom_gaps_before": {}, "bottom_gaps_after": {}}
    assert repair_improves_fit(comparison)
    comparison["zone_utilization_after"]["a"] = 0.83
    assert not repair_improves_fit(comparison)


def test_family_boundary_options_preserve_outer_edges_and_roles():
    zones = [{"zone_id": "top", "x": 0, "y": 0, "width": 1, "height": 0.5, "role": "results"},
             {"zone_id": "bottom", "x": 0, "y": 0.53, "width": 1, "height": 0.47, "role": "supporting"}]
    candidate = {"layout_spec": {"zones": zones}, "exact_bottom_gaps": {"top": 1.5}}
    original = deepcopy(candidate)
    options = boundary_options(candidate, 20)
    assert candidate == original
    assert len(options) == 2
    for _, modified, _, _ in options:
        assert modified[1]["y"] + modified[1]["height"] == pytest.approx(1)
        assert 0.015 - 1e-6 <= modified[1]["y"] - modified[0]["height"] <= 0.03 + 1e-6
        assert [z["role"] for z in modified] == ["results", "supporting"]


def test_partial_blocker_does_not_create_false_full_width_gap():
    layout = [_container("a", "left", 1, 2, 3, 8), _container("b", "center", 5, 2, 3, 7.9),
              _container("c", "right", 9, 2, 2, 8)]
    result = _whitespace(layout)
    assert result["largest_horizontal_gap_inches"] == 1


def test_shared_row_boundary_handles_spanning_modules():
    zones = [{"zone_id": "upper", "x": 0, "y": 0, "width": 1, "height": .5},
             {"zone_id": "lower_a", "x": 0, "y": .53, "width": .48, "height": .47},
             {"zone_id": "lower_b", "x": .52, "y": .53, "width": .48, "height": .47}]
    candidate = {"layout_spec": {"zones": zones}, "exact_bottom_gaps": {"upper": 1.5}}
    options = boundary_options(candidate, 20)
    assert options
    for _, changed, _, _ in options:
        assert changed[1]["y"] == changed[2]["y"]
        assert changed[1]["y"] + changed[1]["height"] == pytest.approx(1)
        assert changed[2]["y"] + changed[2]["height"] == pytest.approx(1)


def test_strict_topology_refinement_is_a_noop():
    candidate = {"layout_spec": {"family_id": "method_flow_spine"}}
    result, audit = refine_composition(object(), candidate, {}, {"template_validation_topology": {"locked": True}}, {}, None, None)
    assert result is candidate and audit["status"] == "locked_or_disabled"


def test_critic_sees_cross_zone_gaps_and_bottom_without_failing_all_underfill():
    layout = [_container("a", "left_top", 1, 7.12, 15, 5),
              _container("b", "left_bottom", 1, 16, 15, 16)]
    zones = [{"zone_id": "left_top", "x": 0, "y": 0, "width": 0.3, "height": 0.25},
             {"zone_id": "left_bottom", "x": 0, "y": 0.32, "width": 0.3, "height": 0.68}]
    state = {"poster_width": 54, "poster_height": 36, "generation_route": "optimization_family",
             "layout_spec": {"family_id": "method_flow_spine", "zones": zones}, "styled_layout": layout}
    result = SectionFitEvaluator().evaluate(CriticContext(state=state, config=load_critic_config()))
    assert any(item["issue_id"].startswith("CROSS-ZONE-GAP") for item in result.issues)
    assert any(item["issue_id"] == "POSTER-BOTTOM-WHITESPACE" for item in result.issues)
    assert state["section_fit_analysis"]["gaps_acceptable"] is False


def test_preflight_rechecks_final_title_font_and_minimum_gap():
    title = _title_layout()[0]
    title.update(measure_loop_title("Results", 4, "Arial", 44, load_config()))
    title.update({"font_family": "Arial", "font_size": 80, "height": 0.8})
    layout = [_container("s", "top", 1, 1, 5, 5), title, _title_layout()[1]]
    issues = validate_intra_section_layout(layout)
    assert any(item["code"] == "PREFLIGHT_FINAL_TITLE_HEIGHT" for item in issues)


def test_preflight_uses_current_font_for_minimum_gap_not_saved_threshold():
    title = _title_layout()[0]
    title.update(measure_loop_title("Results", 4, "Arial", 44, load_config()))
    title.update(font_family="Arial", font_size=80, height=3)
    visual = {**_title_layout()[1], "y": 4.16}
    layout = [_container("s", "top", 1, 1, 5, 7), title, visual]
    issues = validate_intra_section_layout(layout)
    assert any(item["code"] == "PREFLIGHT_TITLE_CONTENT_GAP" for item in issues)


def test_rendered_body_identity_cannot_replace_visual_bounds():
    layout = [_container("s", "top", 1, 2, 4, 7),
              {"type": "visual", "section_id": "s", "x": 1, "y": 2, "width": 4, "height": 5},
              {"type": "text", "section_id": "s", "x": 1, "y": 8, "width": 4, "height": 1}]
    result = measure_composition_whitespace({}, layout, poster_width=12, poster_height=12,
                poster_margin=1, title_height_fraction=.1,
                rendered_boxes=[{"element_id": "s", "rendered_box_inches": [1, 8.1, 5, 8.8]}])
    assert result["mean_terminal_bottom_gap_inches"] == pytest.approx(2.2)
    # The top of the module must still include its image, not start at body text.
    from src.layout.whitespace import visible_module_envelopes
    assert visible_module_envelopes(layout, [{"element_id": "s", "rendered_box_inches": [1, 8.1, 5, 8.8]}])[0]["y"] == 2


def _correction_fixture(tmp_path, monkeypatch, *, second_status="passed", reject=False):
    import src.agents.layout_preflight as preflight
    import src.agents.renderer as renderer
    import src.agents.render_fit_audit as audit_module
    import src.layout.feasibility as feasibility
    import src.layout.families as families
    layout = [_container("s", "top", 1, 1, 4, 4), *_title_layout(),
              _container("next", "top", 1, 7, 4, 2),
              _container("side", "side", 7, 1, 4, 6)]
    title = {"element_id": "s_title", "section_id": "s", "element_type": "section_title",
             "planned_box_inches": [1, 1, 4, 1], "rendered_box_inches": [1.1, 1.2, 4.8, 2.3]}
    problem = {"code": "RENDER_TITLE_CONTENT_GAP", "element_id": "s_title", "section_id": "s",
               "gap_inches": .1, "minimum_gap_inches": .2}
    report = {"status": "failed", "matched_text_boxes": [title], "issues": [problem]}
    state = {"output_dir": str(tmp_path), "poster_width": 12, "poster_height": 12,
             "iteration_context": {"portfolio_mode": "explore"}, "layout_spec": {"zones": []},
             "design_layout": deepcopy(layout), "styled_layout": deepcopy(layout), "errors": []}
    calls = []
    monkeypatch.setattr(feasibility, "validate_compiled_layout", lambda *_args, **_kwargs: {"passed": True})
    def check(trial, *, persist=True):
        assert persist is False
        trial["layout_preflight_failed"] = reject
        trial["layout_preflight_report"] = {"passed": not reject}
        return trial
    monkeypatch.setattr(preflight, "layout_preflight_node", check)
    def render(trial):
        calls.append("render")
        return trial
    monkeypatch.setattr(renderer, "renderer_node", render)
    monkeypatch.setattr(families, "analyze_regions", lambda *_args, **_kwargs: {})
    monkeypatch.setattr(audit_module, "audit_rendered_text", lambda _state: {
        "status": second_status, "passed": True, "issues": [], "issue_count": 0,
        "matched_text_boxes": [title] if second_status == "passed" else [],
        "internal_title_gaps": [{"section_id": "s", "status": "measured"}] if second_status == "passed" else [],
    })
    return state, report, calls


def test_local_title_correction_once_preserves_content_and_other_zones(tmp_path, monkeypatch):
    state, report, calls = _correction_fixture(tmp_path, monkeypatch)
    before = deepcopy(state)
    result = _correct_title_fit(state, report)
    assert result["status"] == "passed" and calls == ["render"]
    for field in ("styled_layout", "design_layout"):
        old, new = before[field], state[field]
        assert new[0]["height"] == pytest.approx(old[0]["height"] + .35)
        assert new[1]["height"] == pytest.approx(old[1]["height"] + .35)
        assert new[2]["y"] == pytest.approx(old[2]["y"] + .35)
        assert new[3]["y"] == pytest.approx(old[3]["y"] + .35)
        assert new[4] == old[4]
        assert new[1]["section_title"] == old[1]["section_title"]
    assert (tmp_path / "content/layout_preflight_report.json").is_file()


def test_rejected_title_correction_does_not_overwrite_original_reports(tmp_path, monkeypatch):
    state, report, calls = _correction_fixture(tmp_path, monkeypatch, reject=True)
    content = tmp_path / "content"
    content.mkdir()
    original_path = content / "layout_preflight_report.json"
    original_path.write_text('{"passed": true, "original": true}')
    before = deepcopy(state)
    assert _correct_title_fit(state, report) is None
    assert calls == [] and state == before
    assert original_path.read_text() == '{"passed": true, "original": true}'
    assert report["title_correction"]["status"] == "contract_rejected"


def test_correction_keeps_compiled_and_rendered_whitespace_separate(tmp_path, monkeypatch):
    import src.agents.render_fit_audit as audit_module
    state, report, calls = _correction_fixture(tmp_path, monkeypatch)
    state["layout_feasibility_report"] = {"status": "feasible", "exact_score": .8}
    original_audit = audit_module.audit_rendered_text
    def rendered(trial):
        result = original_audit(trial)
        result["composition_whitespace"] = {"visible_bounds_source": "rendered_test", "mean_terminal_bottom_gap_inches": 123}
        return result
    monkeypatch.setattr(audit_module, "audit_rendered_text", rendered)
    _correct_title_fit(state, report)
    fit = state["layout_feasibility_report"]
    assert calls == ["render"]
    assert fit["composition_whitespace"]["visible_bounds_source"] == "planned_children"
    assert fit["composition_whitespace"]["mean_terminal_bottom_gap_inches"] != 123
    assert fit["rendered_composition_whitespace"]["mean_terminal_bottom_gap_inches"] == 123
    assert fit["exact_score"] == .8 and fit["exact_score_scope"] == "pre_render_selection"


@pytest.mark.parametrize("status", ["partial", "unavailable"])
def test_inconclusive_correction_cannot_clear_confirmed_failure(tmp_path, monkeypatch, status):
    state, report, calls = _correction_fixture(tmp_path, monkeypatch, second_status=status)
    result = _correct_title_fit(state, report)
    assert calls == ["render"] and result["status"] == "failed" and not result["passed"]
    assert result["issues"][0]["code"] == "RENDER_CORRECTION_UNVERIFIED"


def test_title_correction_rejects_geometry_before_render(tmp_path, monkeypatch):
    import src.layout.feasibility as feasibility
    state, report, calls = _correction_fixture(tmp_path, monkeypatch)
    before = deepcopy(state)
    monkeypatch.setattr(feasibility, "validate_compiled_layout", lambda *_args, **_kwargs: {"passed": False, "issues": [{"code": "PREFLIGHT_ZONE_OVERFLOW"}]})
    assert _correct_title_fit(state, report) is None
    assert state == before and calls == []


def test_title_correction_render_failure_retains_generation_error(tmp_path, monkeypatch):
    import src.agents.renderer as renderer
    state, report, calls = _correction_fixture(tmp_path, monkeypatch)
    def failed(trial):
        calls.append("render")
        trial["errors"].append("renderer: test generation failure")
        trial["attempt_failure_report"] = {"phase": "renderer"}
        return trial
    monkeypatch.setattr(renderer, "renderer_node", failed)
    result = _correct_title_fit(state, report)
    assert calls == ["render"] and not result["passed"]
    assert result["reason"] == "title_correction_render_failed"
    assert state["errors"] == ["renderer: test generation failure"]
    assert state["attempt_failure_report"]["phase"] == "renderer"


def test_refinement_is_strictly_bounded_and_preserves_content(monkeypatch):
    import src.layout.composition_refinement as module
    import src.layout.families as families
    from src.layout.feasibility import _fingerprint
    from types import SimpleNamespace
    monkeypatch.setattr(families, "valid_family_topology", lambda *_: True)
    def options(candidate, *_args, **_kwargs):
        return [(1, [{**candidate["layout_spec"]["zones"][0], "height": .9 - i * .01}], "upper", "lower") for i in range(4)]
    monkeypatch.setattr(module, "boundary_options", options)
    candidate = {"layout_spec": {"family_id": "method_flow_spine", "zones": [{"zone_id": "a", "height": 1}]},
                 "variant": {}, "candidate_content_selection": {"visible_claim_ids": ["c"], "visible_visual_ids": ["v"]},
                 "composition_whitespace": {}, "exact_score": 0, "exact_bottom_gaps": {}}
    calls = []
    def exact(_agent, trial, *_args, **_kwargs):
        calls.append(1)
        trial["exact_score"] += 1
        return trial
    result, audit = refine_composition(SimpleNamespace(poster_margin=1, title_height_fraction=.18), candidate, {},
                                       {"poster_height": 36}, {}, exact, _fingerprint)
    assert len(calls) <= 8 and audit["exact_candidates_checked"] == len(calls)
    assert result["candidate_content_selection"] == candidate["candidate_content_selection"]


@pytest.mark.parametrize("settings", [{"max_rounds": 3}, {"candidates_per_round": 5}, {"enabled": "yes"},
                                      {"preferred_zone_gutter_body_lines": float("nan")}])
def test_invalid_refinement_budget_fails_config_validation(tmp_path, settings):
    import yaml
    config = load_critic_config()
    config["composition_refinement"].update(settings)
    path = tmp_path / "config.yaml"
    path.write_text(yaml.safe_dump(config))
    with pytest.raises(ValueError):
        load_critic_config(path, use_cache=False)


@pytest.mark.parametrize("status,value", [("partial", None), ("measured", True), ("measured", False)])
def test_template_retains_unknown_gap_evidence(status, value):
    from src.templates.library import TemplateDistiller
    state = {"poster_width": 54, "poster_height": 36, "generation_route": "optimization_family",
             "styled_layout": [], "section_fit_analysis": {"gaps_acceptable": value, "gap_check_status": status}}
    template = TemplateDistiller().distill(state, {"candidate_id": "test", "overall_score": 90},
                                          paper_profile={"paper_id": "paper"})
    assert template["density_targets"]["gaps_acceptable"] is value
    assert template["density_targets"]["gap_check_status"] == status

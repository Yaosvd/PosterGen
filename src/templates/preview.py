"""Deterministic, paper-free PNG/PPTX layout previews.

Both outputs are drawn from the same normalized scene. PowerPoint elements are
editable shapes; original claims, figures and exact poster imagery are omitted.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Iterable, List

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from src.layout.families import LayoutFamilyRegistry

PALETTE = {
    "background": (248, 250, 253), "navy": (28, 48, 82),
    "zone": (229, 237, 247), "hero": (207, 225, 246),
    "module": (255, 255, 255), "line": (113, 139, 175),
    "text": (45, 67, 101), "visual": (198, 221, 242),
}


def _box(item: Dict[str, Any]) -> tuple[float, float, float, float]:
    x = max(0.0, min(1.0, float(item.get("x") or 0)))
    y = max(0.0, min(1.0, float(item.get("y") or 0)))
    w = max(0.0, min(1.0 - x, float(item.get("width") or 0)))
    h = max(0.0, min(1.0 - y, float(item.get("height") or 0)))
    return x, y, w, h


def _scene(template: Dict[str, Any]) -> List[Dict[str, Any]]:
    scene: List[Dict[str, Any]] = [{"kind": "banner", "label": "TITLE / AUTHORS", "x": .018, "y": .028, "width": .964, "height": .115}]
    topology = template.get("zone_topology") or {}
    for zone in topology.get("zones") or []:
        x, y, w, h = _box(zone)
        scene.append({"kind": "hero" if zone.get("role") == "hero" else "zone",
                      "label": str(zone.get("zone_id") or "ZONE")[:32],
                      "x": .02 + .96*x, "y": .18 + .79*y,
                      "width": .96*w, "height": .79*h})
    # Accepted templates include actual module placement, not just family zones.
    # Role names are generic; scientific text and figure pixels never appear.
    for item in template.get("normalized_blueprint") or []:
        if item.get("role") not in {"section_container", "visual"}:
            continue
        x, y, w, h = _box(item)
        scene.append({"kind": "visual" if item.get("role") == "visual" else "module",
                      "label": "VISUAL" if item.get("role") == "visual" else "MODULE",
                      "x": x, "y": y, "width": w, "height": h})
    return [item for item in scene if item["width"] > 0 and item["height"] > 0]


def render_preview(template: Dict[str, Any], png_path: Path, pptx_path: Path) -> None:
    """Render one normalized scene to raster and editable PowerPoint."""
    png_path, pptx_path = Path(png_path), Path(pptx_path)
    png_path.parent.mkdir(parents=True, exist_ok=True)
    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    scene = _scene(template)
    aspect = float(template.get("source_aspect_ratio") or 1.6)
    if not 1.0 <= aspect <= 2.5:
        raise ValueError("Invalid template aspect ratio")
    pixels = (1600, round(1600 / aspect))
    canvas = Image.new("RGB", pixels, PALETTE["background"])
    draw = ImageDraw.Draw(canvas)
    bundled_font = Path(__file__).resolve().parents[2] / "fonts" / "Arial.ttf"
    font = ImageFont.truetype(str(bundled_font), 22)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(16), Inches(16 / aspect)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for item in scene:
        x, y, w, h = (item[key] for key in ("x", "y", "width", "height"))
        kind = item["kind"]
        fill = PALETTE["navy" if kind == "banner" else kind if kind in PALETTE else "module"]
        label = item["label"]
        xy = (round(x*pixels[0]), round(y*pixels[1]), round((x+w)*pixels[0]), round((y+h)*pixels[1]))
        draw.rounded_rectangle(xy, radius=10, fill=fill, outline=PALETTE["line"], width=2)
        draw.text((xy[0]+12, xy[1]+12), label, fill=(255,255,255) if kind == "banner" else PALETTE["text"], font=font)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(16*x), Inches(16*y/aspect), Inches(16*w), Inches(16*h/aspect))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*fill)
        shape.line.color.rgb = RGBColor(*PALETTE["line"])
        shape.text = label
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(14 if kind == "banner" else 11)
        paragraph.font.color.rgb = RGBColor(*(255,255,255) if kind == "banner" else PALETTE["text"])
    canvas.save(png_path, format="PNG")
    presentation.save(pptx_path)
    with Image.open(png_path) as check:
        check.verify()
    if len(Presentation(pptx_path).slides[0].shapes) != len(scene):
        raise ValueError("PPTX preview is incomplete")


def generate_family_previews(directory: Path) -> List[Path]:
    """Materialize the twelve registered initial skeletons."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    created: List[Path] = []
    for family in LayoutFamilyRegistry().all():
        topology = {"zones": [dict(zone) for zone in family.zones]}
        template = {"zone_topology": topology}
        png = directory / f"{family.family_id}.png"
        pptx = directory / f"{family.family_id}.pptx"
        render_preview(template, png, pptx)
        created.extend((png, pptx))
    return created


if __name__ == "__main__":
    generate_family_previews(Path(__file__).resolve().parents[2] / "resource" / "layout_families")

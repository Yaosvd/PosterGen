"""Local python-pptx 1.0.2 adapter; no dependency edits or global monkey-patches."""

from importlib.metadata import version
from functools import lru_cache

from PIL import ImageFont
try:
    from pptx.text.layout import TextFitter
except ImportError:  # An unverified future backend may move its private helper.
    TextFitter = None

from src.layout.measurement_context import current_context, file_identity


def _rendered_size(text, point_size, font_file):
    context = current_context()
    identity = file_identity(font_file)
    def compute():
        font = context.memo("font", (identity, point_size),
                            lambda: ImageFont.truetype(font_file, point_size))
        try:
            width, height = font.getsize(text)
        except AttributeError:
            left, top, right, bottom = font.getbbox(text)
            width, height = right - left, bottom - top
        return int(width / 72.0 * 914400), int(height / 72.0 * 914400)
    return context.memo("glyph", (text, point_size, identity), compute)


class OptimizedFitError(RuntimeError):
    """An adapter error must not become a legacy 'text does not fit' result."""


class _CachedTextFitter(TextFitter or tuple):
    def _wrap_lines(self, line_source, point_size):
        context = current_context()
        key = (line_source._text, self._width, point_size, file_identity(self._font_file))
        def compute():
            lines = super(_CachedTextFitter, self)._wrap_lines(line_source, point_size)
            # _wrap_lines returns strings, not the _Line search-tree objects.
            return tuple(lines)
        lines = context.memo("wrap", key, compute)
        return list(lines)

    def _fits_in_width_predicate(self, point_size):
        return lambda line: _rendered_size(line.text, point_size, self._font_file)[0] <= self._width

    @property
    def _fits_inside_predicate(self):
        def predicate(point_size):
            lines = self._wrap_lines(self._line_source, point_size)
            height = _rendered_size("Ty", point_size, self._font_file)[1]
            return height * len(lines) <= self._height
        return predicate


@lru_cache(maxsize=1)
def backend_versions():
    return version("python-pptx"), version("Pillow")


def fit_text(text_frame, font_file, max_size):
    context = current_context()
    if (context is None or not context.settings["enabled"]
            or not context.settings["reuse_wrapping"]):
        return text_frame.fit_text(font_file=font_file, max_size=max_size)
    if backend_versions() != ("1.0.2", "10.4.0"):
        context.fallbacks["unverified_dependency_version"] += 1
        return text_frame.fit_text(font_file=font_file, max_size=max_size)
    if TextFitter is None or not all(hasattr(TextFitter, name) for name in (
        "_wrap_lines", "_fits_in_width_predicate", "best_fit_font_size",
    )):
        context.fallbacks["unverified_backend_interface"] += 1
        return text_frame.fit_text(font_file=font_file, max_size=max_size)
    try:
        file_identity(font_file)
    except OSError:
        context.fallbacks["font_identity_unavailable"] += 1
        return text_frame.fit_text(font_file=font_file, max_size=max_size)
    if text_frame.text == "":
        return None
    # Same input, extents, best-fit search and default formatting as fit_text().
    try:
        size = _CachedTextFitter.best_fit_font_size(
            text_frame.text, text_frame._extents, max_size, font_file,
        )
        text_frame._apply_fit("Calibri", size, False, False)
    except (TypeError, ValueError, OSError) as exc:
        # Preserve genuine old-library narrow-box/font errors, but never hide
        # an adapter regression behind the old measurement's broad handler.
        try:
            text_frame.fit_text(font_file=font_file, max_size=max_size)
        except Exception as original:
            if type(original) is type(exc) and str(original) == str(exc):
                raise original
        raise OptimizedFitError("Cached fitting disagreed with original failure behavior") from exc
    except Exception as exc:
        raise OptimizedFitError("Unexpected cached fitting error") from exc
